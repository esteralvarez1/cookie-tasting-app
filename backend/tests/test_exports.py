"""Tests for export endpoints.

Column mapping (canonical CSV header → meaning):
  ASPECTO:MENCION, ASPECTO:DESCRIPTOR, ASPECTO:VALORACION
  OLOR:MENCION,    OLOR:DESCRIPTOR,    OLOR:VALORACION
  TEXTURA:MENCION, TEXTURA:DESCRIPTOR, TEXTURA:VALORACION
  SABOR:MENCION,   SABOR:DESCRIPTOR,   SABOR:VALORACION
  TEXTO           = full accumulated participant text
  COMPARA         = 1 if any comparison was detected, else 0
  VAGO            = 1 if initial response was vague, else 0
"""
from __future__ import annotations

import csv
import io
import re

import pytest

from app.api.routes.evaluations import service
from app.api.routes.exports import (
    _build_export_filename,
    _iso_export_date,
    _slugify_filename_part,
)
from app.services.exporters import CANONICAL_STRUCTURED_HEADER
from tests.fakes import FakeAnalyzer

_RICH_MESSAGE = (
    'El color es dorado y la forma redonda, la textura es muy crujiente, '
    'sabe dulce con un toque de vainilla y huele muy bien, me gusta mucho'
)


def _complete_evaluation(client, session_id: str, sample_code: str, session_headers: dict) -> None:
    ev = client.post(
        '/api/v1/evaluations',
        json={'session_id': session_id, 'sample_code': sample_code},
        headers=session_headers,
    )
    assert ev.status_code == 201
    eval_id = ev.json()['data']['evaluation_id']
    resp = client.post(
        f'/api/v1/evaluations/{eval_id}/dialog',
        json={'user_message': _RICH_MESSAGE},
        headers=session_headers,
    )
    assert resp.status_code == 200


class TestExportAuth:
    def test_conversations_requires_auth(self, client):
        resp = client.get('/api/v1/export/conversations')
        assert resp.status_code == 401

    def test_user_responses_requires_auth(self, client):
        resp = client.get('/api/v1/export/user-responses')
        assert resp.status_code == 401

    def test_structured_requires_auth(self, client):
        resp = client.get('/api/v1/export/structured')
        assert resp.status_code == 401

    def test_wrong_key_is_rejected(self, client):
        resp = client.get(
            '/api/v1/export/structured',
            headers={'X-Researcher-Key': 'wrong-key'},
        )
        assert resp.status_code == 401

    def test_researcher_key_is_accepted(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200

    def test_admin_key_is_accepted(self, client, admin_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=admin_headers)
        assert resp.status_code == 200


class TestExportRequiresFilter:
    def test_conversations_without_filter_returns_400(self, client, researcher_headers):
        resp = client.get('/api/v1/export/conversations', headers=researcher_headers)
        assert resp.status_code == 400

    def test_user_responses_without_filter_returns_400(self, client, researcher_headers):
        resp = client.get('/api/v1/export/user-responses', headers=researcher_headers)
        assert resp.status_code == 400

    def test_structured_without_filter_returns_400(self, client, researcher_headers):
        resp = client.get('/api/v1/export/structured', headers=researcher_headers)
        assert resp.status_code == 400

    def test_session_id_filter_accepted(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/conversations?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200

    def test_participant_code_filter_accepted(self, client, researcher_headers):
        resp = client.get('/api/v1/export/conversations?participant_code=P001', headers=researcher_headers)
        assert resp.status_code == 200

    def test_tasting_session_config_id_filter_accepted(self, client, researcher_headers, tasting_config):
        config_id = tasting_config['tasting_session_id']
        resp = client.get(f'/api/v1/export/conversations?tasting_session_config_id={config_id}', headers=researcher_headers)
        assert resp.status_code == 200


class TestConversationsExport:
    def test_csv_contains_expected_columns(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/conversations?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        reader = csv.DictReader(io.StringIO(resp.text))
        expected = {
            'participant_code', 'session_id', 'evaluation_id',
            'sample_code', 'turn_index', 'speaker', 'message_text',
        }
        assert expected.issubset(set(reader.fieldnames or []))

    def test_csv_includes_created_turns(self, client, researcher_headers, participant_session, session_headers, monkeypatch):
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.NORMAL))
        session_id = participant_session['session_id']

        ev = client.post(
            '/api/v1/evaluations',
            json={'session_id': session_id, 'sample_code': 'MUESTRA_A'},
            headers=session_headers,
        )
        eval_id = ev.json()['data']['evaluation_id']
        client.post(
            f'/api/v1/evaluations/{eval_id}/dialog',
            json={'user_message': _RICH_MESSAGE},
            headers=session_headers,
        )

        resp = client.get(f'/api/v1/export/conversations?session_id={session_id}', headers=researcher_headers)
        rows = list(csv.DictReader(io.StringIO(resp.text)))
        assert len(rows) > 0
        codes = {row['sample_code'] for row in rows}
        assert 'MUESTRA_A' in codes

    def test_json_format_works(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/conversations?format=json&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        assert resp.json()['success'] is True


class TestStructuredExport:
    def test_csv_has_canonical_header(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        reader = csv.DictReader(io.StringIO(resp.text))
        assert set(reader.fieldnames or []) == set(CANONICAL_STRUCTURED_HEADER)

    def test_csv_includes_completed_evaluation_data(
        self, client, researcher_headers, participant_session, session_headers, monkeypatch
    ):
        """After completing MUESTRA_A, the structured export must have one data row."""
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.ALL_COMPLETE))
        session_id = participant_session['session_id']
        _complete_evaluation(client, session_id, 'MUESTRA_A', session_headers)

        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        rows = list(csv.DictReader(io.StringIO(resp.text)))
        assert len(rows) == 1
        assert rows[0]['TEXTO'] == _RICH_MESSAGE

    def test_modality_columns_populated_after_completion(
        self, client, researcher_headers, participant_session, session_headers, monkeypatch
    ):
        """ALL_COMPLETE fake fills all modality fields; CSV must reflect them."""
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.ALL_COMPLETE))
        session_id = participant_session['session_id']
        _complete_evaluation(client, session_id, 'MUESTRA_A', session_headers)

        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        rows = list(csv.DictReader(io.StringIO(resp.text)))
        row = rows[0]
        for modality in ('ASPECTO', 'OLOR', 'TEXTURA', 'SABOR'):
            # Fake sets these to 'descriptor concreto' etc., so must not be '0' or empty
            assert row[f'{modality}:DESCRIPTOR'] not in ('', '0')
            assert row[f'{modality}:VALORACION'] not in ('', '0')

    def test_empty_modality_fields_are_strings_not_zero(
        self, client, researcher_headers, participant_session, session_headers, admin_headers, monkeypatch
    ):
        """When a modality has no text, exported fields must be '' not '0'."""
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.ALL_COMPLETE))
        session_id = participant_session['session_id']
        _complete_evaluation(client, session_id, 'MUESTRA_A', session_headers)

        # Use admin finalize on a fresh evaluation with NORMAL scenario (no modality data)
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.NORMAL))
        ev = client.post(
            '/api/v1/evaluations',
            json={'session_id': session_id, 'sample_code': 'MUESTRA_B'},
            headers=session_headers,
        )
        assert ev.status_code == 201
        eval_id = ev.json()['data']['evaluation_id']
        client.post(
            f'/api/v1/evaluations/{eval_id}/dialog',
            json={'user_message': 'ok'},
            headers=session_headers,
        )
        client.post(
            f'/api/v1/evaluations/{eval_id}/finalize',
            json={'reason': 'TEST'},
            headers=admin_headers,
        )

        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        rows = list(csv.DictReader(io.StringIO(resp.text)))
        # Find the MUESTRA_B row (the one we force-finalized with no modality data)
        row_b = next((r for r in rows if r.get('Sample_Name') == 'MUESTRA_B'), None)
        if row_b:
            for modality in ('ASPECTO', 'OLOR', 'TEXTURA', 'SABOR'):
                assert row_b[f'{modality}:MENCION'] != '0', f"{modality}:MENCION should be '' not '0'"
                assert row_b[f'{modality}:DESCRIPTOR'] != '0', f"{modality}:DESCRIPTOR should be '' not '0'"
                assert row_b[f'{modality}:VALORACION'] != '0', f"{modality}:VALORACION should be '' not '0'"

    def test_in_progress_evaluation_excluded_from_structured(
        self, client, researcher_headers, participant_session, session_headers, monkeypatch
    ):
        """Structured export only includes COMPLETED evaluations."""
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.NORMAL))
        session_id = participant_session['session_id']
        ev = client.post(
            '/api/v1/evaluations',
            json={'session_id': session_id, 'sample_code': 'MUESTRA_A'},
            headers=session_headers,
        )
        eval_id = ev.json()['data']['evaluation_id']
        # Send one message but do NOT complete the evaluation
        client.post(
            f'/api/v1/evaluations/{eval_id}/dialog',
            json={'user_message': _RICH_MESSAGE},
            headers=session_headers,
        )

        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        rows = list(csv.DictReader(io.StringIO(resp.text)))
        assert len(rows) == 0

    def test_filter_by_tasting_session_config_id(
        self, client, admin_headers, researcher_headers, tasting_config, participant_session, session_headers, monkeypatch
    ):
        """tasting_session_config_id filter must return only data for that config."""
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.ALL_COMPLETE))
        _complete_evaluation(client, participant_session['session_id'], 'MUESTRA_A', session_headers)

        # Create a second tasting config and complete an evaluation in it
        config2_resp = client.post(
            '/api/v1/tasting-sessions',
            json={'title': 'Cata 2', 'sample_codes': ['EXTRA']},
            headers=admin_headers,
        )
        config2_id = config2_resp.json()['data']['tasting_session_id']
        sess2 = client.post(
            '/api/v1/sessions',
            json={'participant_code': 'P999', 'tasting_session_id': config2_id},
        ).json()['data']
        headers2 = {'X-Session-Token': sess2['session_token']}
        _complete_evaluation(client, sess2['session_id'], 'EXTRA', headers2)

        # Filter by the first config — must return exactly 1 row
        config1_id = tasting_config['tasting_session_id']
        resp = client.get(
            f'/api/v1/export/structured?tasting_session_config_id={config1_id}',
            headers=researcher_headers,
        )
        rows = list(csv.DictReader(io.StringIO(resp.text)))
        assert len(rows) == 1

    def test_xlsx_format_returns_bytes(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?format=xlsx&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        assert resp.headers['content-type'].startswith(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        assert len(resp.content) > 0

    def test_json_format_returns_list(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?format=json&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        assert isinstance(resp.json()['data'], list)


class TestStructuredExportCleanColumns:
    _REMOVED_COLUMNS = frozenset({
        'Test_Name', 'Section_Name', 'Section_Number',
        'Panelist_Name', 'Reference_Sample', 'Replicate_Number',
        'INDICE', 'Sample_Set_Number', 'Rep_Number', 'Rep_Name',
        'Sample_Position', 'Sample_Number', 'Design_Position_Name',
        'Blinding_Code', 'Q2__Add_question_name', 'Q3__Add_question_name',
        'Q3__Add_question_name_COMMENTS', 'Q4__Add_question_name',
    })
    _REQUIRED_METADATA = ('Panelist_Code', 'Sample_Name', 'Serving_Order')

    def test_csv_has_required_metadata_columns(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        fieldnames = set(csv.DictReader(io.StringIO(resp.text)).fieldnames or [])
        for col in self._REQUIRED_METADATA:
            assert col in fieldnames, f'{col} must be present in structured CSV export'

    def test_csv_does_not_contain_removed_columns(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        fieldnames = set(csv.DictReader(io.StringIO(resp.text)).fieldnames or [])
        for col in self._REMOVED_COLUMNS:
            assert col not in fieldnames, f'{col} must NOT be present in structured CSV export'

    def test_xlsx_has_required_metadata_columns(self, client, researcher_headers, participant_session):
        import openpyxl
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?format=xlsx&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        headers = [cell.value for cell in ws[1]]
        for col in self._REQUIRED_METADATA:
            assert col in headers, f'{col} must be present in XLSX export'

    def test_xlsx_does_not_contain_removed_columns(self, client, researcher_headers, participant_session):
        import openpyxl
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?format=xlsx&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        headers = [cell.value for cell in ws[1]]
        for col in self._REMOVED_COLUMNS:
            assert col not in headers, f'{col} must NOT be present in XLSX export'

    def test_csv_row_values_after_completion(
        self, client, researcher_headers, participant_session, session_headers, monkeypatch
    ):
        monkeypatch.setattr(service, 'analyzer', FakeAnalyzer(FakeAnalyzer.ALL_COMPLETE))
        session_id = participant_session['session_id']
        _complete_evaluation(client, session_id, 'MUESTRA_A', session_headers)

        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        rows = list(csv.DictReader(io.StringIO(resp.text)))
        assert len(rows) == 1
        row = rows[0]
        assert row['Sample_Name'] == 'MUESTRA_A'
        assert int(row['Serving_Order']) >= 1
        assert row['Panelist_Code']  # non-empty


class TestContentDisposition:
    # The participant_session fixture is linked to the 'Test Cata' config → slug 'test-cata'.
    def test_conversations_csv_has_content_disposition(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/conversations?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        expected = f'conversacion-completa_test-cata_{_iso_export_date()}.csv'
        assert expected in resp.headers.get('content-disposition', '')

    def test_conversations_csv_content_type(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/conversations?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        assert 'text/csv' in resp.headers.get('content-type', '')

    def test_conversations_json_has_content_disposition(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/conversations?format=json&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        expected = f'conversacion-completa_test-cata_{_iso_export_date()}.json'
        assert expected in resp.headers.get('content-disposition', '')

    def test_user_responses_csv_has_content_disposition(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/user-responses?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        expected = f'respuestas-usuario_test-cata_{_iso_export_date()}.csv'
        assert expected in resp.headers.get('content-disposition', '')

    def test_user_responses_csv_content_type(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/user-responses?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        assert 'text/csv' in resp.headers.get('content-type', '')

    def test_structured_csv_has_content_disposition(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        expected = f'clasificacion-estructurada_test-cata_{_iso_export_date()}.csv'
        assert expected in resp.headers.get('content-disposition', '')

    def test_structured_csv_content_type(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        assert 'text/csv' in resp.headers.get('content-type', '')

    def test_structured_json_has_content_disposition(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?format=json&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        expected = f'clasificacion-estructurada_test-cata_{_iso_export_date()}.json'
        assert expected in resp.headers.get('content-disposition', '')

    def test_structured_xlsx_has_content_disposition(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?format=xlsx&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        expected = f'clasificacion-estructurada_test-cata_{_iso_export_date()}.xlsx'
        assert expected in resp.headers.get('content-disposition', '')

    def test_structured_xlsx_content_type(self, client, researcher_headers, participant_session):
        session_id = participant_session['session_id']
        resp = client.get(f'/api/v1/export/structured?format=xlsx&session_id={session_id}', headers=researcher_headers)
        assert resp.status_code == 200
        assert 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' in resp.headers.get('content-type', '')

    def test_content_disposition_is_exposed_via_cors(self, client, researcher_headers, participant_session):
        # The browser must be able to read Content-Disposition on cross-origin requests.
        # CORSMiddleware only emits Access-Control-Expose-Headers when the request carries
        # an allowed Origin (default test origin: http://localhost:5173).
        session_id = participant_session['session_id']
        headers = {**researcher_headers, 'Origin': 'http://localhost:5173'}
        resp = client.get(f'/api/v1/export/structured?session_id={session_id}', headers=headers)
        exposed = resp.headers.get('access-control-expose-headers', '')
        assert 'content-disposition' in exposed.lower()

    def test_filename_falls_back_to_sesion_for_participant_code_only(self, client, researcher_headers):
        # participant_code alone is ambiguous → slug falls back to 'sesion'.
        resp = client.get('/api/v1/export/conversations?participant_code=P001', headers=researcher_headers)
        assert resp.status_code == 200
        expected = f'conversacion-completa_sesion_{_iso_export_date()}.csv'
        assert expected in resp.headers.get('content-disposition', '')


class TestExportFilename:
    def test_slugify_basic_examples(self):
        assert _slugify_filename_part('Test Cata') == 'test-cata'
        assert _slugify_filename_part('Cata Junio 2026') == 'cata-junio-2026'
        assert _slugify_filename_part('Sesión María') == 'sesion-maria'

    def test_slugify_strips_problematic_characters(self):
        assert _slugify_filename_part('A/B:C*?<>|"') == 'abc'

    def test_slugify_collapses_and_trims_hyphens(self):
        assert _slugify_filename_part('  hola   mundo  ') == 'hola-mundo'
        assert _slugify_filename_part('__raro--__') == 'raro'

    def test_slugify_empty_or_none_falls_back(self):
        assert _slugify_filename_part('') == 'sesion'
        assert _slugify_filename_part(None) == 'sesion'  # type: ignore[arg-type]
        assert _slugify_filename_part('   ') == 'sesion'

    def test_iso_export_date_format(self):
        assert re.fullmatch(r'\d{4}-\d{2}-\d{2}', _iso_export_date())

    def test_build_export_filename_full(self):
        date = _iso_export_date()
        assert _build_export_filename('conversacion-completa', 'Prueba 1', 'csv') == f'conversacion-completa_prueba-1_{date}.csv'
        assert _build_export_filename('clasificacion-estructurada', 'Sesión María', 'xlsx') == f'clasificacion-estructurada_sesion-maria_{date}.xlsx'
        assert _build_export_filename('clasificacion-estructurada', 'Test Cata', 'json') == f'clasificacion-estructurada_test-cata_{date}.json'

    def test_build_export_filename_none_title(self):
        date = _iso_export_date()
        assert _build_export_filename('respuestas-usuario', None, 'json') == f'respuestas-usuario_sesion_{date}.json'
